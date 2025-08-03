from flask import Flask, render_template, request, session, redirect
from werkzeug.utils import secure_filename
import os
import PyPDF2
import docx
import openpyxl
import cohere

app = Flask(__name__)
app.secret_key = os.urandom(24)

UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {'pdf', 'docx', 'xlsx'}
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# Cohere setup
co = cohere.Client("uY6A7W64ORC45Jy2RuVxURTnA02gmaO2BEQkOZOa")  
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_file(filepath):
    ext = filepath.rsplit('.', 1)[1].lower()
    text = ""

    if ext == 'pdf':
        with open(filepath, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() or ""

    elif ext == 'docx':
        doc = docx.Document(filepath)
        for para in doc.paragraphs:
            text += para.text + "\n"

    elif ext == 'xlsx':
        wb = openpyxl.load_workbook(filepath)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"

    return text.strip()

@app.route('/', methods=['GET', 'POST'])
def upload_chat():
    if 'chat_history' not in session:
        session['chat_history'] = []

    if request.method == 'POST':
        if 'file' in request.files:
            files = request.files.getlist('file')
            full_text = ""
            for file in files:
                if file and allowed_file(file.filename):
                    filename = secure_filename(file.filename)
                    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
                    file.save(filepath)
                    full_text += extract_text_from_file(filepath) + "\n"

            session['doc_text'] = full_text
            session.modified = True
            return redirect('/')

        elif 'question' in request.form:
            question = request.form['question']
            doc_text = session.get('doc_text', '')
            if not doc_text:
                answer = "Please upload a document first."
            else:
                response = co.chat(
                    message=question,
                    documents=[{"title": "Uploaded Document", "snippet": doc_text[:2000]}]  # Limit to 2k characters
                )
                answer = response.text

            session['chat_history'].append({"sender": "user", "text": question})
            session['chat_history'].append({"sender": "bot", "text": answer})
            session.modified = True
            return redirect('/')

    return render_template('app_upload.html', chat_history=session['chat_history'])

if __name__ == '__main__':
    app.run(debug=True)
